#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
成本核算「离线」自检：用假客户端把整条链路跑通，不需要金蝶账号 / 不需要内网。

覆盖：
    1. 任意日期区间取数（整月 / 5 号~15 号），区间外的记录不进来
    2. 缺采购价子件能被扫描出来（run.missing）
    3. 缺价料号自动预填进维护表 CSV（表头合适、单价列留空、不覆盖已填的价）
    4. 财务在 CSV 里填完价 → 「重新检查」（重算）→ 手工维护料本单列出来，
       自动料本 + 手工维护料本 = 最终料本
    5. 补齐之前 run.complete 为 False（调用方据此卡住导出）；补齐后为 True
    6. 快照存/取 + 基于快照重算（不重新登录金蝶）
    7. 导出的 Excel sheet 与列是否正确（含「区间汇总」合计）
    8. 维护表容错：表头别名、GBK、重复料号、被 Excel 占用时的友好报错
    9. 口径变更拦截（含税/未税、覆盖汇率不一致时不允许套用旧快照）

用法：
    /opt/anaconda3/bin/python3 scripts/verify_costing_offline.py
"""

import csv
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from kingdee_sdk import KingdeeAPIError
from kingdee_sdk.costing import (
    BOM_ENTRY_FIELDS,
    BOM_FORM,
    BOM_HEAD_FIELDS,
    PRICE_FROM_KINGDEE,
    PRICE_FROM_MANUAL,
    PUR_FIELDS,
    PUR_FORM,
    SALE_FIELDS,
    SALE_FORM,
    NO_PRICE_NOTE,
    CostingCalculator,
    CostingError,
    check_recalc_compatible,
    export_xlsx,
    load_snapshot,
    month_bounds,
    month_range,
    normalize_range,
    parse_date_text,
    range_label,
    recalculate,
    save_snapshot,
    summarize_by_product,
)
from kingdee_sdk.manual_prices import CSV_HEADERS, ManualPriceTable, ManualPriceError

FAILURES = []
CHECKS = [0]


def check(label, condition, detail=""):
    CHECKS[0] += 1
    if condition:
        print(f"  [OK]   {label}")
    else:
        FAILURES.append(f"{label} {detail}".strip())
        print(f"  [FAIL] {label}  {detail}")


def check_eq(label, got, want, tol=1e-6):
    if isinstance(want, float):
        ok = isinstance(got, (int, float)) and abs(float(got) - want) <= tol
    else:
        ok = got == want
    check(label, ok, f"（期望 {want!r}，实际 {got!r}）")


# ==================== 假客户端 ====================

class FakeClient:
    """只实现 execute_bill_query，按 costing.py 里用到的过滤/排序语义返回数据"""

    def __init__(self):
        self.sales = [
            # FBillNo, FDate, 状态, 币别, 汇率, 物料, 名称, 数量, 未税单价, 含税单价, FID
            self._sale("SO001", "2026-09-05", "PRE001", 1.0, "P1", "产品一", 1, 100.0, 113.0, 1),
            self._sale("SO002", "2026-09-20", "PRE001", 1.0, "P2", "产品二", 2, 200.0, 226.0, 2),
            self._sale("SO003", "2026-10-03", "PRE001", 1.0, "P3", "产品三", 1, 300.0, 339.0, 3),
            self._sale("SO004", "2026-09-10", "PRE007", 6.7809, "P1", "产品一", 1, 50.0, 56.5, 4),
            self._sale("SO005", "2026-09-25", "PRE001", 1.0, "P9", "无BOM产品", 1, 90.0, 101.7, 5),
        ]
        self.bom_heads = [
            self._bom_head(101, "1.P1"), self._bom_head(102, "1.P2"), self._bom_head(103, "1.P3"),
        ]
        self.bom_entries = [
            self._bom(101, "1.P1", "C1", "材料一", 2.0, 1.0, "Pcs"),
            self._bom(101, "1.P1", "C2", "材料二", 3.0, 1.0, "Pcs"),
            self._bom(102, "1.P2", "C1", "材料一", 1.0, 1.0, "Pcs"),
            self._bom(103, "1.P3", "C3", "材料三", 4.0, 1.0, "Pcs"),
        ]
        self.purchases = [
            self._pur("PO001", "2026-08-01", "C1", 10.0, 11.3, 11),
            self._pur("PO002", "2026-08-05", "C3", 5.0, 5.65, 12),
            # C2 故意没有采购价 → 缺价子件
        ]

    # ---- 造行（键名与 costing.py 请求的 field_keys 完全一致）----
    @staticmethod
    def _sale(bill, date, cur, rate, code, name, qty, price, tax_price, fid):
        return {
            "FBillNo": bill, "FDate": date, "FDocumentStatus": "C",
            "FSettleCurrId.FNumber": cur, "FExchangeRate": rate,
            "FMaterialId.FNumber": code, "FMaterialId.FName": name,
            "FQty": qty, "FPrice": price, "FTaxPrice": tax_price, "FID": fid,
        }

    @staticmethod
    def _bom_head(fid, number):
        return {"FID": fid, "FNumber": number}

    @staticmethod
    def _bom(fid, parent, child, child_name, num, den, unit):
        return {
            "FNumber": f"BOM{fid}", "FMaterialId.FNumber": parent,
            "FMaterialIdChild.FNumber": child, "FMaterialIdChild.FName": child_name,
            "FNumerator": num, "FDenominator": den, "FChildUnitID.FNumber": unit,
            "FID": fid,
        }

    @staticmethod
    def _pur(bill, date, code, price, tax_price, fid):
        return {
            "FBillNo": bill, "FDate": date, "FDocumentStatus": "C",
            "FSettleCurrId.FNumber": "PRE001", "FExchangeRate": 1.0,
            "FMaterialId.FNumber": code, "FPrice": price, "FTaxPrice": tax_price, "FID": fid,
        }

    # ---- 查询 ----
    def execute_bill_query(self, form_id, field_keys, filter_string=None, order_string=None,
                           limit=None, **_kw):
        if form_id == SALE_FORM:
            rows = [r for r in self.sales if self._match_sale(r, filter_string)]
        elif form_id == BOM_FORM and "FID" in (filter_string or "") \
                and "FMaterialId" not in (filter_string or ""):
            fid = int(filter_string.split("FID=")[1].split()[0])
            rows = [r for r in self.bom_entries if r["FID"] == fid]
        elif form_id == BOM_FORM:
            code = filter_string.split("FMaterialId.FNumber='")[1].split("'")[0]
            rows = [r for r in self.bom_heads
                    if r["FNumber"].endswith(code) or code in r["FNumber"]]
        elif form_id == PUR_FORM:
            code = filter_string.split("FMaterialId.FNumber='")[1].split("'")[0]
            price_field = "FPrice" if "FPrice > 0" in filter_string else "FTaxPrice"
            rows = [r for r in self.purchases
                    if r["FMaterialId.FNumber"] == code and (r.get(price_field) or 0) > 0]
        else:
            raise KingdeeAPIError(f"未知表单 {form_id}")

        if order_string and "DESC" in order_string.upper():
            rows = sorted(rows, key=lambda r: r.get("FID", 0), reverse=True)
        if limit:
            rows = rows[:limit]
        keys = [k.strip() for k in field_keys.split(",")]
        return [[r.get(k) for k in keys] for r in rows]

    @staticmethod
    def _match_sale(row, filter_string):
        fs = filter_string or ""
        if "FDocumentStatus='C'" in fs and row["FDocumentStatus"] != "C":
            return False
        if "FDate >=" in fs:
            lo = fs.split("FDate >= '")[1].split("'")[0]
            if str(row["FDate"]) < lo:
                return False
        if "FDate <=" in fs:
            hi = fs.split("FDate <= '")[1].split("'")[0]
            if str(row["FDate"]) > hi:
                return False
        return True


# ==================== 用例 ====================

def case_dates():
    print("\n[1] 日期区间")
    check_eq("month_range('2026-09')", month_range("2026-09"), ("2026-09-01", "2026-09-30"))
    check_eq("month_bounds(0) 本月（以 2026-09-18 为今天）",
             month_bounds(0, __import__("datetime").date(2026, 9, 18)),
             ("2026-09-01", "2026-09-30"))
    check_eq("month_bounds(-1) 上月",
             month_bounds(-1, __import__("datetime").date(2026, 1, 5)),
             ("2025-12-01", "2025-12-31"))
    check_eq("parse '2026/9/5'", parse_date_text("2026/9/5"), "2026-09-05")
    check_eq("parse '20260905'", parse_date_text("20260905"), "2026-09-05")
    check_eq("range_label 整月", range_label("2026-09-01", "2026-09-30"), "2026-09（整月）")
    check_eq("range_label 非整月", range_label("2026-09-05", "2026-09-15"),
             "2026-09-05 ~ 2026-09-15")
    try:
        normalize_range("2026-09-15", "2026-09-05")
        check("起始 > 结束 应报错", False, "（没有报错）")
    except CostingError:
        check("起始 > 结束 会报错", True)
    check_eq("normalize_range 正常", normalize_range("2026-09-01", "2026-09-30"),
             ("2026-09-01", "2026-09-30"))


def case_scan(table, client):
    print("\n[2] 扫描缺价子件（整月）")
    calc = CostingCalculator(client, tax_mode="未税", manual_table=table, log=lambda _t: None)
    run = calc.run("2026-09-01", "2026-09-30")
    check_eq("9 月记录数（10-03 的单不在区间内）", len(run.summary), 4)
    check_eq("缺价子件（C2 有 BOM 但没采购价；P9 是无 BOM 产品被当成采购件）",
             run.missing_codes, ["C2", "P9"])
    check_eq("缺价时 complete=False（导出被卡住）", run.complete, False)
    check_eq("缺价子件涉及产品数", run.missing[0]["涉及产品数"], 1)
    check_eq("缺价子件出现次数", run.missing[0]["出现次数"], 2)
    names = [r["销售单号"] for r in run.summary]
    check_eq("按 FID 倒序", names, ["SO005", "SO004", "SO002", "SO001"])

    t = run.totals()
    # SO001: P1 = C1 2×10 = 20 ；SO002: P2 = C1 1×10 ×2 件 = 20 ；
    # SO004: P1 = 20（美元单只影响售价）；SO005: 无 BOM → 自身当采购件，无采购价 → 0
    check_eq("自动料本合计", t["自动料本"], 60.0)
    check_eq("手工维护料本合计（还没补价）", t["手工维护料本"], 0.0)
    check_eq("最终料本合计", t["最终料本"], 60.0)
    check_eq("销售额合计（100+400+50×6.7809+90）", t["销售额"], 929.045)

    check_eq("每条记录：自动单件 + 手工单件 = 最终单件",
             all(abs(r["自动料本-单件(人民币)"] + r["手工维护料本-单件(人民币)"]
                     - r["最终料本-单件(人民币)"]) < 1e-9 for r in run.summary), True)
    check_eq("缺价明细行备注", 
             sorted({d["备注"] for d in run.details if d["价格来源"] == ""}), [NO_PRICE_NOTE])
    check_eq("金蝶价明细行的价格来源",
             sorted({d["价格来源"] for d in run.details if d["备注"] == ""}), [PRICE_FROM_KINGDEE])
    check_eq("SO005 无 BOM 产品自己成为叶子件且缺价",
             [d["子件编码"] for d in run.details if d["销售单号"] == "SO005"], ["P9"])
    return run


def case_partial_range(client, table):
    print("\n[3] 任意区间：本月 5 号到 15 号")
    calc = CostingCalculator(client, tax_mode="未税", manual_table=table, log=lambda _t: None)
    run = calc.run("2026-09-05", "2026-09-15")
    check_eq("区间内只有 09-05 与 09-10 两条", sorted(r["单据日期"] for r in run.summary),
             ["2026-09-05", "2026-09-10"])
    check_eq("区间标签", run.meta["期间"], "2026-09-05 ~ 2026-09-15")
    check_eq("区间内销售额", run.totals()["销售额"], 439.045)
    return run


def case_csv(table, run):
    print("\n[4] 缺价料号预填进维护表 CSV")
    stats = table.add_missing(run.missing)
    check_eq("新增缺价料号 2 行（C2 与 P9）", stats["added"], 2)
    table.save()
    check("CSV 文件已生成", os.path.exists(table.path), table.path)

    with open(table.path, "r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    check_eq("表头顺序（单价紧跟名称/单位之后，方便直接往后填）",
             rows[0][:5], ["子件编码", "子件名称", "单位", "未税单价", "含税单价"])
    check_eq("预填的料号", rows[1][0], "C2")
    check_eq("预填的名称", rows[1][1], "材料二")
    check_eq("预填的单位", rows[1][2], "Pcs")
    check_eq("单价列留空等财务填", (rows[1][3], rows[1][4]), ("", ""))
    with open(table.path, "rb") as f:
        check_eq("带 BOM 的 utf-8-sig（Excel 双击不乱码）", f.read(3), b"\xef\xbb\xbf")

    # 再扫一次：不应重复新增，也不能动已填的价
    table.rows["C2"]["未税单价"] = "4"
    stats2 = table.add_missing(run.missing)
    check_eq("重复扫描不新增行", stats2["added"], 0)
    check_eq("重复扫描不覆盖已填的价", table.rows["C2"]["未税单价"], "4")
    table.rows["C2"]["未税单价"] = ""
    table.save()


def fill_csv_like_excel(table, prices):
    """模拟财务用 Excel 打开 CSV、在对应行填价后保存（没有的行就像手工新增一行）"""
    table.load()
    for code, values in prices.items():
        row = table.rows.get(code)
        if row is None:
            row = {h: "" for h in CSV_HEADERS}
            row["子件编码"] = code
            table.rows[code] = row
        for col, val in values.items():
            row[col] = val
    table.save()
    return table


def case_fill_and_recheck(run, table, tmp):
    print("\n[5] 财务填价后「重新检查」→ 手工维护料本单列")
    fill_csv_like_excel(table, {
        "C2": {"未税单价": "4", "含税单价": "4.52", "币别": "人民币", "维护人": "财务小张"},
        "P9": {"未税单价": "30", "含税单价": "33.9", "币别": "人民币", "维护人": "财务小张"},
    })
    table2 = ManualPriceTable(table.path).load()
    check_eq("读回维护表行数", len(table2.rows), 2)
    check("lookup 命中 C2", table2.lookup("C2", "未税")["ok"])
    check("lookup 命中 P9", table2.lookup("P9", "未税")["ok"])

    run2 = recalculate(run, table2, override_rate=None)
    check_eq("补齐后 complete=True（允许导出）", run2.complete, True)
    check_eq("缺价清单清空", run2.missing, [])

    t = run2.totals()
    # C2 用量 3 × 4 = 12/件；SO001 与 SO004 都是 P1 → 12×2 = 24
    # P9 无 BOM，自身 1 件 × 30 = 30
    check_eq("手工维护料本合计", t["手工维护料本"], 54.0)
    check_eq("自动料本合计（金蝶价部分不变）", t["自动料本"], 60.0)
    check_eq("最终料本合计 = 自动 + 手工", t["最终料本"], 114.0)
    check_eq("毛利合计 = 销售额 - 最终料本", t["毛利"], 929.045 - 114.0)

    row = [r for r in run2.summary if r["销售单号"] == "SO001"][0]
    check_eq("SO001 自动单件", row["自动料本-单件(人民币)"], 20.0)
    check_eq("SO001 手工维护单件（单列出来）", row["手工维护料本-单件(人民币)"], 12.0)
    check_eq("SO001 最终单件 = 20 + 12", row["最终料本-单件(人民币)"], 32.0)
    check_eq("SO001 手工维护子件数", row["手工维护子件数"], 1)
    check("SO001 备注标出手工维护", "手工维护价" in row["备注"], row["备注"])

    manual_details = [d for d in run2.details if d["价格来源"] == PRICE_FROM_MANUAL]
    check_eq("明细里手工维护行数（C2 两条 + P9 一条）", len(manual_details), 3)
    c2 = [d for d in manual_details if d["销售单号"] == "SO001" and d["子件编码"] == "C2"][0]
    check_eq("手工维护行单价", c2["采购价(人民币)"], 4.0)
    check_eq("手工维护行金额（用量 3 × 4）", c2["金额(人民币)"], 12.0)
    check_eq("手工维护行没有采购单号", c2["采购单号"], "")
    check_eq("手工维护行的价格来源列", c2["价格来源"], "手工维护")
    return run2, table2


def case_snapshot(run, table2, run2, tmp):
    print("\n[6] 快照：补价不用重新登录金蝶")
    snap = os.path.join(tmp, "snap.json.gz")
    save_snapshot(run, snap)
    check("快照已写出", os.path.exists(snap))
    loaded = load_snapshot(snap)
    check_eq("快照记录数", len(loaded.records), 4)
    check_eq("快照区间", loaded.meta["期间"], "2026-09（整月）")
    check_eq("快照回读后仍是缺价状态", loaded.complete, False)
    again = recalculate(loaded, table2, override_rate=None)
    check_eq("基于快照重算 = 直接重算",
             [r["最终料本(人民币)"] for r in again.summary],
             [r["最终料本(人民币)"] for r in run2.summary])
    check_eq("基于快照重算的手工料本",
             again.totals()["手工维护料本"], run2.totals()["手工维护料本"])

    print("\n[6b] 口径变更拦截")
    check("含税/未税不一致 → 拦住", 
          check_recalc_compatible(loaded.meta, "含税", None) is not None)
    check("覆盖汇率不一致 → 拦住",
          check_recalc_compatible(loaded.meta, "未税", 6.9) is not None)
    check("口径一致 → 放行", check_recalc_compatible(loaded.meta, "未税", None) is None)


def case_stale(client, table2):
    print("\n[7] 维护表里有价、但金蝶已有价 → 提示未采用")
    fill_csv_like_excel(table2, {"C1": {"未税单价": "8"}})
    calc = CostingCalculator(client, tax_mode="未税", manual_table=table2, log=lambda _t: None)
    run = calc.run("2026-09-01", "2026-09-30")
    check_eq("金蝶价优先（C1 出现在 3 条记录里，都按 10 算）",
             [d["采购价(人民币)"] for d in run.details
              if d["子件编码"] == "C1" and d["价格来源"] == PRICE_FROM_KINGDEE],
             [10.0, 10.0, 10.0])
    check_eq("提示 C1 未采用手工价", [s["子件编码"] for s in run.stale], ["C1"])
    check_eq("C2 / P9 已填 → 无缺价", run.missing_codes, [])
    check_eq("手工维护料本 = C2 24 + P9 30", run.totals()["手工维护料本"], 54.0)
    # 清掉 C1，避免影响后续断言
    table2.rows.pop("C1")
    table2.save()


def case_partial_fill(run, table2):
    print("\n[8] 只填了一部分 → 仍然卡住导出")
    fill_csv_like_excel(table2, {"C2": {"未税单价": ""}, "P9": {"未税单价": ""}})
    run3 = recalculate(run, ManualPriceTable(table2.path).load())
    check_eq("清空后又变成缺价", run3.complete, False)
    check_eq("两个都缺（C2 出现多、排前）", run3.missing_codes, ["C2", "P9"])
    c2 = [m for m in run3.missing if m["子件编码"] == "C2"][0]
    check("缺价原因提示可读", "维护表" in c2["说明"], c2["说明"])

    fill_csv_like_excel(table2, {"C2": {"未税单价": "4.5", "含税单价": "5.085"}})
    run4 = recalculate(run, ManualPriceTable(table2.path).load())
    check_eq("只填了 C2、P9 还空着 → 依旧拦住导出", run4.complete, False)
    check_eq("缺价清单只剩 P9", run4.missing_codes, ["P9"])
    check_eq("已填的 C2 立刻生效（3×4.5=13.5/件）",
             [r["手工维护料本-单件(人民币)"] for r in run4.summary
              if r["销售单号"] == "SO001"], [13.5])

    fill_csv_like_excel(table2, {"P9": {"未税单价": "30"}})
    run5 = recalculate(run, ManualPriceTable(table2.path).load())
    check_eq("全部填完 → 放行导出", run5.complete, True)
    # 还原成用例 5/9 用的价，保证后续断言一致
    fill_csv_like_excel(table2, {"C2": {"未税单价": "4", "含税单价": "4.52"}})


def case_export(run2, tmp):
    print("\n[9] 导出 Excel")
    out = os.path.join(tmp, "costing.xlsx")
    export_xlsx(out, run2, extra_meta={"测试": "离线自检"})
    check("xlsx 已生成", os.path.exists(out))

    from openpyxl import load_workbook
    wb = load_workbook(out, read_only=True)
    check_eq("工作表", wb.sheetnames,
             ["逐条核算", "区间汇总", "子件价格总览", "子件明细", "缺价待维护", "说明"])
    ws = wb["逐条核算"]
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    for col in ("自动料本-单件(人民币)", "手工维护料本-单件(人民币)", "最终料本-单件(人民币)",
                "自动料本(人民币)", "手工维护料本(人民币)", "最终料本(人民币)",
                "毛利(人民币)", "毛利率"):
        check(f"逐条核算含列「{col}」", col in header)

    ws_all = wb["子件价格总览"]
    all_rows = list(ws_all.iter_rows(values_only=True))
    ia = {n: i for i, n in enumerate(all_rows[0])}
    check_eq("子件价格总览列", list(all_rows[0]),
             ["子件编码", "子件名称", "单位", "价格来源", "采购价-未税(原币)", "单价(人民币)",
              "出现次数", "涉及产品数"])
    check_eq("总览列出全部子件（C1/C2/P9）",
             sorted(r[ia["子件编码"]] for r in all_rows[1:]), ["C1", "C2", "P9"])
    src = {r[ia["子件编码"]]: r[ia["价格来源"]] for r in all_rows[1:]}
    check_eq("总览里的价格来源", src, {"C1": "金蝶", "C2": "手工维护", "P9": "手工维护"})

    ws2 = wb["区间汇总"]
    rows2 = list(ws2.iter_rows(values_only=True))
    idx = {n: i for i, n in enumerate(rows2[0])}
    check_eq("区间汇总行数（3 个产品 + 合计）", len(rows2) - 1, 4)
    check_eq("合计行标签", rows2[-1][idx["产品编码"]], "合计")
    check_eq("合计最终料本", rows2[-1][idx["最终料本(人民币)"]], 114.0)
    check_eq("合计手工维护料本", rows2[-1][idx["手工维护料本(人民币)"]], 54.0)
    check_eq("合计销售额", rows2[-1][idx["销售额(人民币)"]], 929.045)

    ws3 = wb["缺价待维护"]
    rows3 = list(ws3.iter_rows(values_only=True))
    check_eq("缺价待维护提示（已补齐）", rows3[1][0], "（无）")
    return out


def case_tolerance(tmp):
    print("\n[10] 维护表容错")
    p = os.path.join(tmp, "alias.csv")
    with open(p, "w", encoding="gbk", newline="") as f:
        w = csv.writer(f)
        w.writerow(["料号", "名称", "单位", "未税价", "含税价", "币种"])
        w.writerow(["C2", "材料二", "Pcs", "6.5", "7.345", "美元"])
        w.writerow(["C2", "材料二(重复)", "Pcs", "9", "9", "人民币"])
        w.writerow(["C9", "材料九", "Pcs", "", "", ""])
    t = ManualPriceTable(p).load()
    check_eq("GBK + 表头别名能认出来", sorted(t.rows), ["C2", "C9"])
    check_eq("重复料号取第一条并记录告警", t.duplicates, ["C2"])
    res = t.lookup("C2", "未税")
    check_eq("别名列取到价", res["price"], 6.5)
    check_eq("币别名（美元）保留原文，汇率没填则为 0", (res["currency"], res["rate"]), ("美元", 0.0))
    r2 = t.lookup("C9", "未税")
    check_eq("没填价 → 不可用", r2["ok"], False)
    check("没填价的原因可读", "还没填" in r2["reason"], r2["reason"])

    p2 = os.path.join(tmp, "bad.csv")
    with open(p2, "w", encoding="utf-8-sig", newline="") as f:
        f.write("完全不对的表头,第二列\n1,2\n")
    t2 = ManualPriceTable(p2).load()
    check_eq("表头认不出 → 不报错、视为空表", (len(t2.rows), t2.header_ok), (0, False))

    # 被 Excel 占用：把目标路径换成目录，制造不可写
    p3 = os.path.join(tmp, "locked.csv")
    t3 = ManualPriceTable(p3)
    t3.rows["X1"] = {h: "" for h in CSV_HEADERS}
    t3.rows["X1"]["子件编码"] = "X1"
    os.makedirs(p3, exist_ok=True)   # 同名目录 → 写文件必然失败
    try:
        t3.save()
        check("写不进去时应给出友好提示", False, "（竟然成功了）")
    except ManualPriceError as exc:
        check("写不进去时给出友好提示", "Excel" in str(exc), str(exc)[:60])
    shutil.rmtree(p3, ignore_errors=True)


def case_all_kingdee(tmp):
    print("\n[11] 全都有价 → 直接可导出")
    client = FakeClient()
    client.purchases.append(FakeClient._pur("PO003", "2026-08-09", "C2", 7.0, 7.91, 13))
    client.purchases.append(FakeClient._pur("PO004", "2026-08-11", "P9", 40.0, 45.2, 14))
    t = ManualPriceTable(os.path.join(tmp, "empty.csv")).load()
    calc = CostingCalculator(client, tax_mode="含税", manual_table=t, log=lambda _t: None)
    run = calc.run("2026-09-01", "2026-09-30")
    check_eq("含税口径无缺价", run.complete, True)
    # 含税：C1 2×11.3 = 22.6，C2 3×7.91 = 23.73 → 46.33
    check_eq("SO001 最终单件（含税）",
             [r["最终料本-单件(人民币)"] for r in run.summary
              if r["销售单号"] == "SO001"], [46.33])
    check_eq("含税列名", any(k.startswith("售价-含税") for k in run.summary[0]), True)
    check_eq("含税明细列名", any(k.startswith("采购价-含税") for k in run.details[0]), True)


def main():
    tmp = tempfile.mkdtemp(prefix="costing_offline_")
    try:
        print(f"离线自检开始（临时目录 {tmp}）")
        client = FakeClient()
        table = ManualPriceTable(os.path.join(tmp, "采购价维护表.csv")).load()

        case_dates()
        run = case_scan(table, client)
        case_partial_range(client, table)
        case_csv(table, run)
        run2, table2 = case_fill_and_recheck(run, table, tmp)
        case_snapshot(run, table2, run2, tmp)
        case_stale(client, table2)
        case_partial_fill(run, table2)
        case_export(run2, tmp)
        case_tolerance(tmp)
        case_all_kingdee(tmp)
    finally:
        if os.getenv("KEEP_OFFLINE_TMP"):
            print(f"\n[保留临时目录] {tmp}")
        else:
            shutil.rmtree(tmp, ignore_errors=True)

    print("\n" + "=" * 62)
    if FAILURES:
        print(f"离线自检未通过：{CHECKS[0] - len(FAILURES)}/{CHECKS[0]} 项通过，失败 {len(FAILURES)} 项")
        for f in FAILURES:
            print("  ✗ " + f)
        return 1
    print(f"离线自检全部通过：{CHECKS[0]}/{CHECKS[0]} 项 ✅")
    return 0


if __name__ == "__main__":
    sys.exit(main())
