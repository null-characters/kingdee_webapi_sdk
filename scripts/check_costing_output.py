#!/usr/bin/env python3
"""
核对核算输出（读导出的 xlsx）：
    1. 按「销售单号 + 产品」检查明细有没有重复计数
    2. 看 BOM 展开层级是否正常
    3. 统计缺价子件
    4. 统计叶子件采购价的日期分布（提醒「不限时点取最新价」可能用到很久以前的老价）

用法：
    python3 scripts/check_costing_output.py [xlsx路径]
"""

import os
import sys
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

DEFAULT = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "data", "sample_export", "costing_2026-09_未税.xlsx",
)
OLD_PRICE_BEFORE = "2026-01-01"  # 早于这个日期的采购价算「老价」


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT
    if not os.path.exists(path):
        print(f"[找不到文件] {path}")
        return 1

    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    print(f"[读取] {path}")
    print(f"工作表：{wb.sheetnames}")

    ws = wb["子件明细"]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    idx = {name: i for i, name in enumerate(header)}

    groups = {}
    dates = []
    old_prices = []
    for r in rows:
        key = (r[idx["销售单号"]], r[idx["产品编码"]])
        groups.setdefault(key, []).append(r)
        d = r[idx["采购日期"]]
        if d:
            dates.append(str(d))
            if str(d) < OLD_PRICE_BEFORE:
                old_prices.append((str(d), r[idx["子件编码"]], r[idx["子件名称"]], r[idx["采购单号"]]))

    print(f"\n{'销售单号':<17}{'产品编码':<18}{'明细行':>7}{'去重子件':>9}{'重复行':>7}{'缺价':>5}  层级分布")
    total_dup = 0
    for (bill, code), items in groups.items():
        codes = [x[idx["子件编码"]] for x in items]
        key = [(x[idx["子件编码"]], x[idx["层级"]], x[idx["累计用量"]]) for x in items]
        dup = len(key) - len(set(key))
        total_dup += dup
        levels = Counter(x[idx["层级"]] for x in items)
        missing = sum(1 for x in items if x[idx["备注"]] == "缺采购价")
        level_txt = " ".join(f"L{k}:{v}" for k, v in sorted(levels.items()))
        print(f"{bill:<17}{code:<18}{len(items):>7}{len(set(codes)):>9}{dup:>7}{missing:>5}  {level_txt}")
    print(f"\n完全重复行合计：{total_dup}（应为 0，>0 说明同一层同一子件被算了两遍）")

    if dates:
        years = Counter(d[:4] for d in dates)
        print(f"\n叶子件采购价日期分布（共 {len(dates)} 行）：")
        for y, c in sorted(years.items()):
            print(f"    {y} 年：{c} 行")
        if old_prices:
            print(f"\n⚠ 采购价早于 {OLD_PRICE_BEFORE} 的叶子件 {len(old_prices)} 行（老价），最老的 8 条：")
            for d, code, name, bill in sorted(old_prices)[:8]:
                print(f"    {d}  {code:<18}{str(name)[:16]:<18}来自 {bill}")
        else:
            print(f"\n没有早于 {OLD_PRICE_BEFORE} 的老价。")

    ws2 = wb["逐条核算"]
    rows2 = list(ws2.iter_rows(values_only=True))
    print(f"\n逐条核算：{len(rows2) - 1} 条记录")
    sample = rows2[0]
    print("  列：" + " | ".join(str(v) for v in sample))
    return 0


if __name__ == "__main__":
    sys.exit(main())
